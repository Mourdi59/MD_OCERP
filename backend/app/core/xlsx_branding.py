# DDC-CWICR-OE: DataDrivenConstruction - OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""The company letterhead on generated Excel workbooks.

The PDF side prints the firm's letterhead through :mod:`app.core.pdf_branding`.
A bill of quantities, an RFI log or an invoice register that leaves the office
as a spreadsheet is the same kind of document, and a firm that cannot send a
PDF without its letterhead does not want the workbook to go out bare either.

:func:`apply_company_header` is called by an exporter on a finished worksheet,
just before the workbook is saved. It inserts rows above the table and writes
the letterhead into them: the document logo (else the app logo) at the top
left, the registered name beside it, then the address, the registration line
and one contact line, a thin rule, and the document title. Everything below
moves down, and everything that points at a row moves with it: freeze panes,
autofilter, print titles and print area, merged cells, row heights,
conditional formats, data validations, tables, images, hyperlinks and formula
references. So an exporter is wired with one call and none of its own row
arithmetic changes.

The sheet's printed header and footer carry the company name and the page
number, with the header left off the first page, where the letterhead already
is.

**Nothing changes without a letterhead.** The PDF layer makes the decision, so
the PDF and the workbook cannot disagree: a legal name or a document logo in
the company profile, and the appearance switch ``show_letterhead`` on.
Otherwise the function returns ``1`` without touching the sheet, and the
workbook is identical cell for cell to what the exporter wrote.

**Mind the importer.** The rows above the table move the header off row 1.
The BOQ importer looks for its header under a letterhead for that reason
(:func:`app.modules.boq.importers.excel.locate_header_row`); an exporter
whose file is meant to be imported again must not be wired until its
importer does the same.

A sheet holding charts or pivot tables gets the print header and footer but no
rows: their data references are not shifted here, and a chart plotting the
wrong cells is worse than a sheet without a letterhead.

Never raises, like the PDF layer: a logo that cannot be decoded is left out and
the text still prints, and any other failure leaves the export as it was. That
last promise is kept by rolling back: the cells move first and everything that
names a row follows them, so a failure part way through would otherwise ship a
sheet with the table pushed down and nothing above it. Each step records the
inverse of what it changes (:class:`_Undo`), and the failure path runs them.

User text that only looks like a formula, a note reading ``=) done``, is what
made that matter. An exporter types its text as text
(:func:`app.core.xlsx_text.store_strings_as_text`) so the question does not
arise; a cell that reaches here as a formula and cannot be parsed is left
exactly as it is.
"""

from __future__ import annotations

import logging
import re
from collections.abc import Callable
from io import BytesIO
from typing import Any

from app.core import pdf_branding

logger = logging.getLogger(__name__)

#: The logo box, in screen pixels (Excel sizes pictures at 96 dpi). A wide
#: wordmark and a square mark both fit, and the box stays shorter than the four
#: text lines beside it, so the letterhead is no taller with a logo than
#: without one.
_LOGO_MAX_W_PX = 180
_LOGO_MAX_H_PX = 60
#: A logo is stored at up to twice its display size so it prints sharp without
#: carrying a 1200 px SVG raster into every workbook.
_LOGO_STORE_SCALE = 2
#: Beside the logo, the text starts in the first column clear of it. A gap
#: between the two is how letterheads look (the PDF sets the block against
#: the opposite margin); text off screen when the workbook opens is not. So
#: when that column starts further right than this, the text goes under the
#: logo instead. 480 px is the left half of a small laptop's sheet area.
_LOGO_TEXT_GAP_PX = 8
_TEXT_MAX_LEFT_PX = 480

#: Row heights, in points.
_NAME_ROW_PT = 21.0
_DETAIL_ROW_PT = 14.0
_TITLE_ROW_PT = 18.0
_SPACER_ROW_PT = 8.0

#: Same greys as the PDF letterhead.
_MUTED = "666666"
_RULE = "CCCCCC"
_DEFAULT_ACCENT = "1A1A2E"

#: Excel refuses a header or footer longer than 255 characters in total.
_HEADER_PART_CHARS = 100

#: A title passed by the exporter, capped like the PDF brand line.
_TITLE_CHARS = 160

_HEX_COLOUR = re.compile(r"#?([0-9A-Fa-f]{6})")


def apply_company_header(ws: Any, *, title: str | None = None, subtitle: str | None = None) -> int:
    """Put the company letterhead above the table on ``ws``.

    Call it once, on the finished sheet, right before the workbook is saved.

    Args:
        ws: An openpyxl worksheet. A write-only sheet is left alone.
        title: The document title printed under the letterhead, e.g. the
            sheet's own name. ``None`` for a sheet that already opens with
            its title.
        subtitle: A second, lighter line under the title.

    Returns:
        The row the table's first row now sits on: ``1`` when there is no
        letterhead, otherwise the first row below it.
    """
    undo = _Undo()
    inserted = 0
    try:
        profile = _letterhead_profile()
        if profile is None or not hasattr(ws, "insert_rows"):
            return 1
        brand = _brand_name()
        title_text = _single_line(title)
        subtitle_text = _single_line(subtitle)
        if getattr(ws, "_charts", None) or getattr(ws, "_pivots", None):
            _set_print_header_footer(ws, brand, title_text, letterhead=False, undo=undo)
            return 1

        logo = _logo_picture(profile)
        lines = _company_lines(profile)
        if logo is None and not lines:
            # A logo-only profile whose logo will not decode: nothing of the
            # firm's to print, which is where the PDF letterhead stops too.
            _set_print_header_footer(ws, brand, title_text, letterhead=False, undo=undo)
            return 1
        layout = _layout(ws, logo, lines, title_text, subtitle_text)

        inserted = layout.rows
        _remember_styles(ws, undo)
        _move_cells_down(ws, inserted, undo)
        _shift_everything_below(ws, inserted, undo)
        _write_letterhead(ws, layout, logo, lines, title_text, subtitle_text, undo)
        _set_print_header_footer(ws, brand, title_text, letterhead=True, undo=undo)
        return inserted + 1
    except Exception:  # noqa: BLE001 - a letterhead must never break an export
        # Put back what was already changed. A sheet with the table pushed
        # down and no letterhead over it is worse than one without a
        # letterhead, and it is what the exporter would ship otherwise.
        undo.restore()
        logger.warning("Excel letterhead skipped (build failed)", exc_info=True)
        return 1


# -- What to print -------------------------------------------------------------


class _Undo:
    """How to put back everything a letterhead changes, newest change first.

    The promise at the top of this module is that a failure leaves the sheet
    as the exporter wrote it. Moving the cells down is the first change and
    the rest follow it, so a failure part way through used to ship a sheet
    with the table pushed down and nothing over it. Every step registers the
    inverse of what it is about to do, and the failure path runs them.
    """

    def __init__(self) -> None:
        self._steps: list[Callable[[], None]] = []

    def add(self, restore: Callable[[], None]) -> None:
        """Remember how to undo a change about to be made."""
        self._steps.append(restore)

    def attribute(self, obj: Any, name: str) -> None:
        """Remember ``obj.name`` as it is now, before it is set to something else."""
        old = getattr(obj, name)
        self._steps.append(lambda: setattr(obj, name, old))

    def restore(self) -> None:
        """Undo every remembered change, latest first."""
        for step in reversed(self._steps):
            try:
                step()
            except Exception:  # noqa: BLE001 - one step failing must not stop the rest
                logger.warning("Excel letterhead rollback step failed", exc_info=True)
        self._steps.clear()


#: The workbook-wide tables a style lands in when it is given to a cell. The
#: letterhead's own fonts and borders go in them, and a cell that is taken
#: back off the sheet does not take its style out of them.
_STYLE_TABLES = ("_fonts", "_fills", "_borders", "_alignments", "_protections", "_number_formats", "_cell_styles")


def _remember_styles(ws: Any, undo: _Undo) -> None:
    """Remember the workbook's style tables, so a rollback leaves no trace."""
    workbook = getattr(ws, "parent", None)
    if workbook is None:
        return
    for name in _STYLE_TABLES:
        table = getattr(workbook, name, None)
        if table is None:
            continue
        before = list(table)
        undo.add(lambda name=name, before=before, table=table: setattr(workbook, name, type(table)(before)))


def _letterhead_profile() -> dict[str, Any] | None:
    """The company profile when a letterhead is due, else ``None``.

    The PDF layer's two conditions, asked of the PDF layer: the profile calls
    for one and the appearance has not switched it off.
    """
    profile = pdf_branding._letterhead_profile()
    if profile is None:
        return None
    if not pdf_branding.branded_appearance().get("show_letterhead", True):
        return None
    return profile


def _brand_name() -> str:
    """The name for the printed header and footer, or ``""``.

    The same name the PDF footer prints. A workspace with a logo and no name
    anywhere gets no text rather than the platform's name.
    """
    name = pdf_branding._company_name()
    return "" if name == pdf_branding.DEFAULT_BRAND else name


def _accent() -> str:
    colour = pdf_branding.branded_appearance().get("accent_color")
    match = _HEX_COLOUR.fullmatch(colour) if isinstance(colour, str) else None
    return match.group(1).upper() if match else _DEFAULT_ACCENT


def _text(profile: dict[str, Any], field: str) -> str:
    value = profile.get(field)
    return value.strip() if isinstance(value, str) else ""


def _single_line(value: Any) -> str:
    if not isinstance(value, str):
        return ""
    return " ".join(value.split())[:_TITLE_CHARS]


def _company_lines(profile: dict[str, Any]) -> list[tuple[str, str]]:
    """``(kind, text)`` for each line of the company block that has text.

    The address runs on one line, the way the PDF sets it under a centred
    logo: three address rows would add three rows to every frozen pane.
    """
    lines: list[tuple[str, str]] = []
    legal_name = _text(profile, "legal_name")
    if legal_name:
        lines.append(("name", legal_name))
    address = " · ".join(line.strip() for line in _text(profile, "address").split("\n") if line.strip())
    registration = _text(profile, "registration_line")
    contact = " · ".join(part for part in (_text(profile, key) for key in ("phone", "email", "website")) if part)
    lines.extend(("detail", text) for text in (address, registration, contact) if text)
    return lines


class _Picture:
    """A decoded logo: PNG bytes plus the size to show it at, in pixels."""

    def __init__(self, png: bytes, width: int, height: int) -> None:
        self.png = png
        self.width = width
        self.height = height


def _logo_picture(profile: dict[str, Any]) -> _Picture | None:
    """The best logo that decodes, fitted to the logo box, or ``None``.

    Document logo first, then the app logo, exactly as the PDF picks them; an
    SVG arrives already rasterised by :mod:`app.core.pdf_branding`. Everything
    is re-encoded as PNG: openpyxl names the stored file after the source
    format, and Excel will not open a picture part called ``.webp``.
    """
    try:
        from PIL import Image as PILImage
    except ImportError:
        return None
    for data_url in pdf_branding._logo_candidates(pdf_branding._read_branding(), profile):
        try:
            raw = pdf_branding._logo_image_bytes(data_url)
            if not raw:
                continue
            with PILImage.open(BytesIO(raw)) as source:
                source.load()
                image = source.convert("RGBA")
            width, height = image.size
            if not width or not height:
                continue
            # Never enlarged, like the PDF logo: a small raster stretched to
            # fill the box prints soft.
            scale = min(_LOGO_MAX_W_PX / width, _LOGO_MAX_H_PX / height, 1.0)
            shown_w = max(1, round(width * scale))
            shown_h = max(1, round(height * scale))
            image.thumbnail((shown_w * _LOGO_STORE_SCALE, shown_h * _LOGO_STORE_SCALE))
            out = BytesIO()
            image.save(out, format="PNG")
            return _Picture(out.getvalue(), shown_w, shown_h)
        except Exception:  # noqa: BLE001 - try the next logo, never raise
            logger.debug("Could not decode a logo for the Excel letterhead", exc_info=True)
    return None


# -- Layout --------------------------------------------------------------------


class _Layout:
    """Where each part of the letterhead goes, decided before the sheet moves."""

    def __init__(self) -> None:
        self.rows = 0
        #: Row of the logo (1-based), when there is one.
        self.logo_row = 1
        #: Rows of the company lines, and the column they are written in.
        self.line_rows: list[int] = []
        self.text_col = 1
        #: Height, in points, of every letterhead row that is not the default.
        self.heights: dict[int, float] = {}
        self.rule_row = 0
        self.rule_cols = 1
        self.title_row = 0
        self.subtitle_row = 0


def _column_px(ws: Any, col: int) -> int:
    """Width of column ``col`` on screen, in pixels.

    ``column_dimensions.get`` rather than indexing: the dimension holder
    creates an entry for any column it is asked about, and a created entry is
    written out with a width of its own.
    """
    from openpyxl.utils import get_column_letter

    dim = ws.column_dimensions.get(get_column_letter(col))
    width = getattr(dim, "width", None) if dim is not None and getattr(dim, "customWidth", False) else None
    chars = float(width) if width else 8.43
    return int(chars * 7 + 5)


def _layout(
    ws: Any,
    logo: _Picture | None,
    lines: list[tuple[str, str]],
    title: str,
    subtitle: str,
) -> _Layout:
    """Plan the letterhead rows against the finished sheet's column widths.

    The logo sits at the top left and the text beside it, from the first
    column clear of the logo. When the first columns are so wide that the text
    would start out of sight, the logo gets a row of its own and the text goes
    under it from column A instead.
    """
    layout = _Layout()
    table_cols = max(int(ws.max_column or 1), 1)
    heights = [(_NAME_ROW_PT if kind == "name" else _DETAIL_ROW_PT) for kind, _ in lines]
    row = 1

    side_by_side = False
    if logo is not None and lines:
        # The first column whose left edge clears the logo.
        left = 0
        col = 1
        while left < logo.width + _LOGO_TEXT_GAP_PX and left <= _TEXT_MAX_LEFT_PX:
            left += _column_px(ws, col)
            col += 1
        side_by_side = left <= _TEXT_MAX_LEFT_PX
        if side_by_side:
            layout.text_col = col

    logo_pt = logo.height * 0.75 if logo is not None else 0.0
    if logo is not None and not side_by_side:
        # The logo on a row of its own, the text under it from column A.
        layout.logo_row = row
        layout.heights[row] = max(logo_pt + 4.0, _DETAIL_ROW_PT)
        row += 1
    elif logo is not None:
        # Beside the text: the text rows grow to the logo's height if needed.
        shortfall = logo_pt + 4.0 - sum(heights)
        if shortfall > 0:
            heights[-1] += shortfall

    for height in heights:
        layout.line_rows.append(row)
        layout.heights[row] = height
        row += 1

    layout.rule_row = row - 1
    layout.rule_cols = max(table_cols, layout.text_col)
    if title:
        layout.title_row = row
        layout.heights[row] = _TITLE_ROW_PT
        row += 1
    if subtitle:
        layout.subtitle_row = row
        layout.heights[row] = _DETAIL_ROW_PT
        row += 1
    # One short blank row between the letterhead and the table.
    layout.heights[row] = _SPACER_ROW_PT
    layout.rows = row
    return layout


def _write_letterhead(
    ws: Any,
    layout: _Layout,
    logo: _Picture | None,
    lines: list[tuple[str, str]],
    title: str,
    subtitle: str,
    undo: _Undo,
) -> None:
    from openpyxl.styles import Alignment, Border, Font, Side

    for row, height in layout.heights.items():
        # Nothing of the exporter's is left up here: every row it wrote moved
        # down, heights and all, so the letterhead's own rows simply go again.
        undo.add(lambda row=row: ws.row_dimensions.pop(row, None))
        ws.row_dimensions[row].height = height

    name_font = Font(bold=True, size=14, color=_accent())
    detail_font = Font(size=9, color=_MUTED)
    middle = Alignment(vertical="center")
    for (kind, text), row in zip(lines, layout.line_rows, strict=True):
        cell = _text_cell(ws, row, layout.text_col, text)
        cell.font = name_font if kind == "name" else detail_font
        cell.alignment = middle

    if layout.rule_row:
        rule = Border(bottom=Side(style="thin", color=_RULE))
        for col in range(1, layout.rule_cols + 1):
            ws.cell(row=layout.rule_row, column=col).border = rule
    if layout.title_row:
        cell = _text_cell(ws, layout.title_row, 1, title)
        cell.font = Font(bold=True, size=12)
        cell.alignment = Alignment(vertical="bottom")
    if layout.subtitle_row:
        cell = _text_cell(ws, layout.subtitle_row, 1, subtitle)
        cell.font = Font(size=10, color=_MUTED)

    if logo is not None:
        from openpyxl.drawing.image import Image

        picture = Image(BytesIO(logo.png))
        picture.width = logo.width
        picture.height = logo.height
        picture.anchor = f"A{layout.logo_row}"
        ws.add_image(picture)
        # The cells the letterhead writes are cleared by the move's own undo,
        # which owns the rows above the table; the image is not a cell.
        undo.add(lambda: _remove_image(ws, picture))


def _remove_image(ws: Any, picture: Any) -> None:
    """Take the logo back off the sheet."""
    images = getattr(ws, "_images", None)
    if images is not None and picture in images:
        images.remove(picture)


def _text_cell(ws: Any, row: int, col: int, text: str) -> Any:
    """Write ``text`` as a text cell, never a formula.

    Typed as text rather than passed through
    :func:`app.core.csv_safety.neutralise_formula`: that prefixes an
    apostrophe to anything starting with ``+``, and an international phone
    number does. The firm's line must print as the firm wrote it, and a cell
    stored as text is not evaluated when the workbook opens, whatever it
    starts with.

    Control characters are dropped first: openpyxl refuses them, and the
    title is user text arriving after the table has already moved down.
    """
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    cell = ws.cell(row=row, column=col, value=ILLEGAL_CHARACTERS_RE.sub("", text))
    cell.data_type = "s"
    return cell


def _set_print_header_footer(ws: Any, brand: str, title: str, *, letterhead: bool, undo: _Undo) -> None:
    """Company name in the printed header and footer, page number bottom right.

    Left alone when the exporter set a header or footer of its own. With a
    letterhead on the sheet the first page has no header: the letterhead is
    the header there.
    """
    parts = (ws.oddHeader.left, ws.oddHeader.center, ws.oddHeader.right, ws.oddFooter.left, ws.oddFooter.right)
    if any(part.text for part in parts) or ws.oddFooter.center.text:
        return

    def _code(text: str) -> str:
        # ``&`` introduces a header code (``&P`` is the page number), so a
        # literal one is doubled.
        return text[:_HEADER_PART_CHARS].replace("&", "&&")

    if brand:
        undo.attribute(ws.oddHeader.left, "text")
        undo.attribute(ws.oddFooter.left, "text")
        ws.oddHeader.left.text = _code(brand)
        ws.oddFooter.left.text = _code(brand)
    if title:
        undo.attribute(ws.oddHeader.right, "text")
        ws.oddHeader.right.text = _code(title)
    undo.attribute(ws.oddFooter.right, "text")
    ws.oddFooter.right.text = "&P / &N"
    if letterhead:
        undo.attribute(ws.HeaderFooter, "differentFirst")
        undo.attribute(ws.firstFooter.left, "text")
        undo.attribute(ws.firstFooter.right, "text")
        ws.HeaderFooter.differentFirst = True
        ws.firstFooter.left.text = ws.oddFooter.left.text
        ws.firstFooter.right.text = ws.oddFooter.right.text


# -- Moving what points at a row ------------------------------------------------

_CELL_PART = re.compile(r"(\$?[A-Za-z]{1,3})(\$?)(\d+)")
_ROW_PART = re.compile(r"(\$?)(\d+)")
_COL_PART = re.compile(r"\$?[A-Za-z]{1,3}")


def _shift_part(part: str, rows: int) -> str | None:
    """One side of a reference moved down, or ``None`` if it is not a row reference."""
    match = _CELL_PART.fullmatch(part)
    if match:
        return f"{match.group(1)}{match.group(2)}{int(match.group(3)) + rows}"
    match = _ROW_PART.fullmatch(part)
    if match:
        return f"{match.group(1)}{int(match.group(2)) + rows}"
    return None


def _shift_reference(value: str, rows: int, sheet_title: str, *, local: bool) -> str:
    """A reference moved down ``rows`` rows when it points at the moved sheet.

    ``local`` says whether an unqualified reference means the moved sheet,
    true for a formula written on it. Anything that is not plainly a cell,
    cell range or row range (a defined name, a table reference, a whole
    column) comes back as it was.
    """
    sheet, sep, ref = value.rpartition("!")
    if sep:
        name = sheet[1:-1].replace("''", "'") if sheet.startswith("'") and sheet.endswith("'") else sheet
        if name != sheet_title:
            return value
    elif not local:
        return value
    parts = ref.split(":")
    if len(parts) > 2:
        return value
    if all(_COL_PART.fullmatch(part) for part in parts):
        return value
    if len(parts) == 1 and not _CELL_PART.fullmatch(parts[0]):
        return value
    shifted = [_shift_part(part, rows) for part in parts]
    if any(part is None for part in shifted):
        return value
    return f"{sheet}{sep}{':'.join(shifted)}"  # type: ignore[arg-type]


def _shift_formula(formula: Any, rows: int, sheet_title: str, *, local: bool) -> Any:
    """A formula with every reference to the moved sheet moved down.

    Text the tokeniser refuses comes back as it is. A cell typed as a formula
    only because its text opens with ``=``, a note reading ``=) done``, is not
    a formula and names no cell, so there is nothing in it to move; failing
    the whole letterhead over it would help nobody. An exporter should type
    such text as text (:func:`app.core.xlsx_text.store_strings_as_text`), and
    this is what happens to the ones that have not.
    """
    if not (isinstance(formula, str) and formula.startswith("=")):
        return formula
    from openpyxl.formula.tokenizer import Token, Tokenizer

    try:
        tokens = Tokenizer(formula)
        items = list(tokens.items)
    except Exception:  # noqa: BLE001 - openpyxl raises bare exceptions from the tokeniser
        # Only the reading is forgiven. Moving a reference that was read is
        # this module's own work, and a failure there rolls the letterhead
        # back rather than shipping a formula pointing at the wrong rows.
        logger.debug("Excel letterhead left a cell that is not a formula alone", exc_info=True)
        return formula
    changed = False
    for token in items:
        if token.type == Token.OPERAND and token.subtype == Token.RANGE:
            moved = _shift_reference(token.value, rows, sheet_title, local=local)
            if moved != token.value:
                token.value = moved
                changed = True
    return tokens.render() if changed else formula


def _shift_range(value: str, rows: int) -> str:
    """``A1:B2`` (no sheet name) moved down."""
    from openpyxl.worksheet.cell_range import CellRange

    cell_range = CellRange(value)
    cell_range.shift(row_shift=rows)
    return cell_range.coord


def _shift_multi_range(value: Any, rows: int) -> Any:
    from openpyxl.worksheet.cell_range import MultiCellRange

    return MultiCellRange([_shift_range(str(part), rows) for part in value.ranges])


def _move_cells_down(ws: Any, rows: int, undo: _Undo) -> None:
    """Move every cell down ``rows`` rows: ``ws.insert_rows(1, rows)``, faster.

    openpyxl's ``insert_rows`` first creates a cell for every empty coordinate
    in the sheet's used range and then sorts every key, which took about two
    seconds on a 20 000-row, 15-column sheet. Every row moves by the same
    amount here, so the cells are simply re-keyed in one pass. Falls back to
    ``insert_rows`` if openpyxl ever stops keeping its cells in that dict.

    The undo moves them back, and clears the rows the letterhead is about to
    be written into: nothing of the exporter's lives there once it has moved.
    """
    cells = getattr(ws, "_cells", None)
    if not isinstance(cells, dict):
        undo.add(lambda: ws.delete_rows(1, rows))
        ws.insert_rows(1, rows)
        return
    undo.add(lambda: _move_cells_up(ws, rows))
    moved = [(row + rows, col, cell) for (row, col), cell in cells.items()]
    cells.clear()
    for row, col, cell in moved:
        cell.row = row
        cells[(row, col)] = cell
    # Where ``ws.append`` writes next, as ``insert_rows`` leaves it.
    ws._current_row = max((row for row, _col, _cell in moved), default=0)


def _move_cells_up(ws: Any, rows: int) -> None:
    """Undo :func:`_move_cells_down`, letterhead and all.

    Whatever sits in the top ``rows`` rows is the letterhead's own: the
    exporter's cells all moved below it.
    """
    cells = getattr(ws, "_cells", None)
    if not isinstance(cells, dict):
        return
    moved = [(row - rows, col, cell) for (row, col), cell in cells.items() if row > rows]
    cells.clear()
    for row, col, cell in moved:
        cell.row = row
        cells[(row, col)] = cell
    ws._current_row = max((row for row, _col, _cell in moved), default=0)


def _shift_everything_below(ws: Any, rows: int, undo: _Undo) -> None:
    """Move everything that names a row by ``rows``, after the cells moved.

    Moving cells moves nothing else; openpyxl documents that merged cells,
    formulas, row heights and the rest are the caller's. Every step hands
    ``undo`` the inverse of what it is about to change first: a formula this
    sheet cannot tokenise raises here, half way through.
    """
    from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
    from openpyxl.worksheet.formula import ArrayFormula

    title = ws.title

    # Row heights, outline levels, hidden rows.
    moved = sorted(ws.row_dimensions.items(), reverse=True)
    undo.add(lambda: _restore_row_dimensions(ws, moved, rows))
    for index, _ in moved:
        del ws.row_dimensions[index]
    for index, dim in moved:
        dim.index = index + rows
        ws.row_dimensions[index + rows] = dim

    # The ranges live in a set, and shifting one changes its hash, so the set
    # is rebuilt rather than mutated in place.
    merged = list(ws.merged_cells.ranges)
    if merged:
        before = [str(cell_range) for cell_range in merged]
        undo.add(lambda: _restore_merged(ws, before))
    for cell_range in merged:
        cell_range.shift(row_shift=rows)
    ws.merged_cells.ranges = set(merged)

    # A pane that freezes rows keeps freezing the same ones, which now sit
    # under the letterhead; a pane that freezes only columns stays as it is.
    pane = ws.freeze_panes
    if pane:
        row, col = coordinate_to_tuple(pane)
        if row > 1:
            undo.attribute(ws, "freeze_panes")
            ws.freeze_panes = f"{get_column_letter(col)}{row + rows}"

    if ws.auto_filter.ref:
        undo.attribute(ws.auto_filter, "ref")
        ws.auto_filter.ref = _shift_range(ws.auto_filter.ref, rows)
        sort_state = ws.auto_filter.sortState
        if sort_state is not None and sort_state.ref:
            undo.attribute(sort_state, "ref")
            sort_state.ref = _shift_range(str(sort_state.ref), rows)

    if ws.print_title_rows:
        undo.attribute(ws, "print_title_rows")
        first, _, last = ws.print_title_rows.replace("$", "").partition(":")
        ws.print_title_rows = f"{int(first) + rows}:{int(last or first) + rows}"
    if ws.print_area:
        from openpyxl.worksheet.print_settings import PrintArea

        undo.attribute(ws, "print_area")
        area = PrintArea.from_string(ws.print_area)
        ws.print_area = [_shift_range(str(part), rows) for part in area.ranges]

    for brk in ws.row_breaks.brk:
        undo.attribute(brk, "id")
        brk.id += rows

    for table in ws.tables.values():
        undo.attribute(table, "ref")
        table.ref = _shift_range(table.ref, rows)
        if table.autoFilter is not None and table.autoFilter.ref:
            undo.attribute(table.autoFilter, "ref")
            table.autoFilter.ref = _shift_range(table.autoFilter.ref, rows)

    for validation in ws.data_validations.dataValidation:
        for field in ("sqref", "formula1", "formula2"):
            undo.attribute(validation, field)
        validation.sqref = _shift_multi_range(validation.sqref, rows)
        validation.formula1 = _shift_formula_text(validation.formula1, rows, title)
        validation.formula2 = _shift_formula_text(validation.formula2, rows, title)

    if ws.conditional_formatting:
        from openpyxl.formatting.formatting import ConditionalFormattingList

        old = ws.conditional_formatting
        undo.attribute(ws, "conditional_formatting")
        ws.conditional_formatting = ConditionalFormattingList()
        for formatting in old:
            target = " ".join(_shift_range(str(part), rows) for part in formatting.sqref.ranges)
            for rule in formatting.rules:
                undo.attribute(rule, "formula")
                rule.formula = [_shift_formula_text(text, rows, title) for text in rule.formula or []]
                ws.conditional_formatting.add(target, rule)

    for picture in ws._images:
        _shift_anchor(picture, rows, undo)

    # A hyperlink records the cell it sits on when it is set, and the writer
    # reads that record, not the cell's new place. Formulas: on this sheet an
    # unqualified reference means this sheet; on the others only one that
    # names it does.
    workbook = ws.parent
    sheets = workbook.worksheets if workbook is not None else [ws]
    for sheet in sheets:
        local = sheet is ws
        for cell in getattr(sheet, "_cells", {}).values():
            if local:
                link = getattr(cell, "hyperlink", None)
                if link is not None:
                    undo.attribute(link, "ref")
                    link.ref = cell.coordinate
            if getattr(cell, "data_type", None) != "f":
                continue
            value = cell.value
            if isinstance(value, ArrayFormula):
                undo.attribute(value, "text")
                value.text = _shift_formula(value.text, rows, title, local=local)
                if local and value.ref:
                    undo.attribute(value, "ref")
                    value.ref = _shift_range(value.ref, rows)
            else:
                undo.attribute(cell, "value")
                cell.value = _shift_formula(value, rows, title, local=local)
    if workbook is not None:
        for scope in [workbook.defined_names, *(sheet.defined_names for sheet in sheets)]:
            for defined in scope.values():
                if defined.attr_text:
                    undo.attribute(defined, "attr_text")
                    defined.attr_text = _shift_formula("=" + defined.attr_text, rows, title, local=False)[1:]


def _restore_row_dimensions(ws: Any, moved: list[tuple[int, Any]], rows: int) -> None:
    """Put the row heights back where they were before the shift."""
    for index, _dim in moved:
        ws.row_dimensions.pop(index + rows, None)
    for index, dim in moved:
        dim.index = index
        ws.row_dimensions[index] = dim


def _restore_merged(ws: Any, before: list[str]) -> None:
    """Put the merged ranges back, rebuilt from the coordinates they had."""
    from openpyxl.worksheet.cell_range import CellRange

    ws.merged_cells.ranges = {CellRange(coordinates) for coordinates in before}


def _shift_formula_text(text: Any, rows: int, sheet_title: str) -> Any:
    """A validation or conditional-format formula, stored without its ``=``."""
    if not isinstance(text, str) or not text:
        return text
    return _shift_formula("=" + text, rows, sheet_title, local=True)[1:]


def _shift_anchor(picture: Any, rows: int, undo: _Undo) -> None:
    anchor = picture.anchor
    if isinstance(anchor, str):
        from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

        row, col = coordinate_to_tuple(anchor)
        undo.attribute(picture, "anchor")
        picture.anchor = f"{get_column_letter(col)}{row + rows}"
        return
    marker = getattr(anchor, "_from", None)
    if marker is not None:
        undo.attribute(marker, "row")
        marker.row += rows
    end = getattr(anchor, "to", None)
    if end is not None:
        undo.attribute(end, "row")
        end.row += rows


__all__ = ["apply_company_header"]
