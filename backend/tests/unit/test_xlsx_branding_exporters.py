# DDC-CWICR-OE: DataDrivenConstruction - OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""The company letterhead on the second batch of wired Excel exports.

Each exporter is run for real, DB-free: route handlers are called directly
with a stub session or service and the project-access check stubbed out.
For every one of them:

* without a company profile the workbook is byte for byte what it was before
  the letterhead existed, proven by running it again with the letterhead
  call replaced by a no-op and comparing every part of the package;
* with a profile the firm's name heads the sheet and the table's header row
  is intact under it.

The profile is set through the PDF layer's readers, where the letterhead
decision is made, so nothing here depends on a profile saved on this machine.
"""

from __future__ import annotations

import asyncio
import io
import uuid
import zipfile
from collections.abc import Callable
from datetime import UTC, date, datetime
from types import SimpleNamespace
from typing import Any

import pytest
from openpyxl import load_workbook

from app.core import pdf_branding, xlsx_branding
from app.core.company_profile import DEFAULT_COMPANY_PROFILE

_LEGAL_NAME = "Acme & Sons Construction GmbH"
_PROJECT = uuid.UUID("22222222-2222-2222-2222-222222222222")


@pytest.fixture
def use_profile(monkeypatch: pytest.MonkeyPatch):
    """Set the company profile; the default is an empty workspace."""

    def _set(profile: dict[str, str] | None = None) -> None:
        full = dict(DEFAULT_COMPANY_PROFILE)
        full.update(profile or {})
        monkeypatch.setattr(pdf_branding, "_read_company_profile", lambda *_args, **_kw: dict(full))
        monkeypatch.setattr(pdf_branding, "_read_branding", lambda *_args, **_kw: {})
        monkeypatch.setattr(pdf_branding, "_read_appearance", lambda *_args, **_kw: {})

    _set()
    return _set


async def _allowed(*_args: Any, **_kwargs: Any) -> None:
    return None


class _Session:
    """Answers ``await session.execute(stmt)`` with fixed rows."""

    def __init__(self, rows: list[Any]) -> None:
        self._rows = rows

    async def execute(self, _stmt: Any) -> Any:
        rows = self._rows
        return SimpleNamespace(scalars=lambda: SimpleNamespace(all=lambda: list(rows)))


def _body(response: Any) -> bytes:
    async def _collect() -> bytes:
        chunks = [chunk async for chunk in response.body_iterator]
        return b"".join(chunk if isinstance(chunk, bytes) else chunk.encode() for chunk in chunks)

    return asyncio.run(_collect())


# -- The exporters, each as a zero-argument producer of workbook bytes ----------


def _punch_list(monkeypatch: pytest.MonkeyPatch) -> Callable[[], bytes]:
    from app.modules.punchlist.service import PunchListService

    items = [
        SimpleNamespace(
            title="Touch up paint, stair core B",
            status="open",
            priority="high",
            category="finishes",
            trade="painting",
            assigned_to="Painter Ltd",
            due_date=date(2026, 10, 1),
            description="Scuffs on landing 3",
            resolution_notes=None,
            created_at=datetime(2026, 9, 1, 8, 0, tzinfo=UTC),
        )
    ]

    async def _all_for_project(_project_id: Any) -> list[Any]:
        return items

    async def _names(_values: Any) -> dict[str, str]:
        return {}

    service = PunchListService.__new__(PunchListService)
    service.repo = SimpleNamespace(all_for_project=_all_for_project)
    service.resolve_party_names = _names  # type: ignore[method-assign]
    return lambda: asyncio.run(service.export_excel(_PROJECT))


def _inspections(monkeypatch: pytest.MonkeyPatch) -> Callable[[], bytes]:
    from app.modules.inspections import router as inspections_router

    async def _names(_session: Any, _ids: Any) -> dict[str, str]:
        return {}

    monkeypatch.setattr(inspections_router, "verify_project_access", _allowed)
    monkeypatch.setattr(inspections_router, "resolve_party_names", _names)
    rows = [
        SimpleNamespace(
            inspection_number="QI-001",
            title="Rebar check, slab L2",
            inspection_type="structural",
            inspector_id=None,
            inspection_date="2026-09-10",
            location="Level 2",
            status="completed",
            result="pass",
            checklist_data=[{"response": "yes"}, {"response": "no"}],
        )
    ]
    return lambda: _body(asyncio.run(inspections_router.export_inspections(_PROJECT, _Session(rows), "user")))


def _safety_incidents(monkeypatch: pytest.MonkeyPatch) -> Callable[[], bytes]:
    from app.modules.safety import router as safety_router

    monkeypatch.setattr(safety_router, "verify_project_access", _allowed)
    rows = [
        SimpleNamespace(
            incident_number="INC-001",
            incident_date="2026-09-03",
            incident_type="near_miss",
            location="Gate 2",
            description="Unsecured load on crane hook",
            severity="medium",
            treatment_type=None,
            days_lost=0,
            root_cause=None,
            status="open",
            reported_to_regulator=False,
        )
    ]
    return lambda: _body(asyncio.run(safety_router.export_incidents(_PROJECT, _Session(rows), "user", None)))


def _safety_observations(monkeypatch: pytest.MonkeyPatch) -> Callable[[], bytes]:
    from app.modules.safety import router as safety_router

    monkeypatch.setattr(safety_router, "verify_project_access", _allowed)
    rows = [
        SimpleNamespace(
            observation_number="OBS-001",
            created_at=datetime(2026, 9, 4, 9, 30, tzinfo=UTC),
            observation_type="unsafe_condition",
            location="Scaffold north",
            description="Missing toe board",
            severity=3,
            likelihood=2,
            risk_score=6,
            status="open",
            corrective_action="Fit toe board",
        )
    ]
    return lambda: _body(asyncio.run(safety_router.export_observations(_PROJECT, _Session(rows), "user", None)))


def _approvals_register(monkeypatch: pytest.MonkeyPatch) -> Callable[[], bytes]:
    from app.modules.file_approvals import router as approvals_router

    monkeypatch.setattr(approvals_router, "_require_project_access", _allowed)
    workflows = [
        SimpleNamespace(
            file_kind="document",
            file_id="A-101 General arrangement",
            file_version_snapshot="rev C",
            submitted_by_id=None,
            submitted_at=datetime(2026, 9, 5, 12, 0, tzinfo=UTC),
            status="approved",
            final_decision_at=datetime(2026, 9, 6, 12, 0, tzinfo=UTC),
            final_decision_by_id=None,
            notes="Approved as noted",
            steps=[],
        )
    ]

    async def _list_workflows(_project_id: Any) -> list[Any]:
        return workflows

    service = SimpleNamespace(list_workflows=_list_workflows)
    return lambda: _body(
        asyncio.run(approvals_router.export_approvals_register("user", None, service, project_id=_PROJECT))
    )


def _bim_boq(monkeypatch: pytest.MonkeyPatch) -> Callable[[], bytes]:
    from app.modules.bim_hub.exporters.boq_xlsx import BoqExportOptions, build_boq_workbook

    model = SimpleNamespace(id=uuid.uuid4(), name="Tower A")
    elements = [
        SimpleNamespace(
            id=uuid.uuid4(),
            stable_id="wall-a",
            element_type="IfcWall",
            name="Wall A",
            storey="L1",
            discipline="Architectural",
            quantities={"area_m2": 10.0, "volume_m3": 2.0},
        )
    ]
    options = BoqExportOptions(frozen_now=datetime(2026, 9, 22, 12, 0, tzinfo=UTC))
    return lambda: build_boq_workbook(model, elements, options)


def _validation_report(monkeypatch: pytest.MonkeyPatch) -> Callable[[], bytes]:
    from app.core.validation.engine import RuleCategory, RuleResult, Severity, ValidationReport
    from app.modules.validation.tabular_exporter import report_to_xlsx

    report = ValidationReport(
        target_type="boq",
        target_id="boq-123",
        rule_sets_applied=["din276"],
        results=[
            RuleResult(
                rule_id="din276.cost_group",
                rule_name="Cost group",
                severity=Severity.ERROR,
                category=RuleCategory.COMPLIANCE,
                passed=False,
                message="Position 01.002 has no cost group",
            )
        ],
        duration_ms=1.0,
    )
    return lambda: report_to_xlsx(report)


#: (producer, sheet, the table's first header cell)
_EXPORTERS = {
    "punch_list": (_punch_list, "Punch List", "No."),
    "inspections": (_inspections, "Inspections", "Inspection #"),
    "safety_incidents": (_safety_incidents, "Safety Incidents", "Incident #"),
    "safety_observations": (_safety_observations, "Safety Observations", "Observation #"),
    "approvals_register": (_approvals_register, "Approval Register", "File Kind"),
    "bim_boq": (_bim_boq, "BOQ", "Bill of Quantities - Tower A"),
    "validation_report": (_validation_report, "Validation", "Validation findings"),
}


def _parts(blob: bytes) -> dict[str, bytes]:
    """Every part of the package but the one carrying save timestamps."""
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        return {name: archive.read(name) for name in archive.namelist() if name != "docProps/core.xml"}


@pytest.mark.parametrize("name", sorted(_EXPORTERS))
def test_without_a_profile_the_export_is_byte_identical(
    use_profile, monkeypatch: pytest.MonkeyPatch, name: str
) -> None:
    make, sheet, _first = _EXPORTERS[name]
    produce = make(monkeypatch)
    wired = produce()

    calls: list[str] = []

    def _no_letterhead(ws: Any, **_kwargs: Any) -> int:
        calls.append(ws.title)
        return 1

    monkeypatch.setattr(xlsx_branding, "apply_company_header", _no_letterhead)
    unwired = produce()

    # The exporter really calls the helper, on the right sheet, once: a
    # comparison with a patch that never took effect would prove nothing.
    assert calls == [sheet]
    assert _parts(wired) == _parts(unwired)


@pytest.mark.parametrize("name", sorted(_EXPORTERS))
def test_with_a_profile_the_export_carries_the_letterhead(
    use_profile, monkeypatch: pytest.MonkeyPatch, name: str
) -> None:
    make, sheet, first = _EXPORTERS[name]
    use_profile({"legal_name": _LEGAL_NAME, "phone": "+49 30 1234567"})

    ws = load_workbook(io.BytesIO(make(monkeypatch)()))[sheet]

    assert ws["A1"].value == _LEGAL_NAME
    assert ws["A2"].value == "+49 30 1234567"
    table_row = next(row for row in range(1, 30) if ws.cell(row=row, column=1).value == first)
    assert table_row > 3
    assert ws.oddFooter.left.text == "Acme && Sons Construction GmbH"
