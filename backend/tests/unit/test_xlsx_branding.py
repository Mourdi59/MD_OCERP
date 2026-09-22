# DDC-CWICR-OE: DataDrivenConstruction - OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""The company letterhead on Excel exports (``app.core.xlsx_branding``).

Every test sets the company profile, the app branding and the document
appearance through the PDF layer's three readers, which is where the
letterhead decision is made. None of them reads the data dir, so a profile
another session saved on this machine cannot turn a test green or red.
"""

from __future__ import annotations

import base64
import io
import zipfile
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from PIL import Image as PILImage

from app.core import pdf_branding, xlsx_branding
from app.core.company_profile import DEFAULT_COMPANY_PROFILE
from app.core.xlsx_branding import apply_company_header

_PROFILE = {
    "legal_name": "Acme & Sons Construction GmbH",
    "address": "Hauptstrasse 1\n10115 Berlin\nGermany",
    "registration_line": "HRB 12345 · VAT DE123456789",
    "phone": "+49 30 1234567",
    "email": "office@acme.example",
    "website": "acme.example",
}


def _png_data_url(width: int = 400, height: int = 100, fmt: str = "PNG") -> str:
    image = PILImage.new("RGB", (width, height), (20, 60, 140))
    out = io.BytesIO()
    image.save(out, format=fmt)
    mime = {"PNG": "png", "WEBP": "webp", "JPEG": "jpeg"}[fmt]
    return f"data:image/{mime};base64," + base64.b64encode(out.getvalue()).decode("ascii")


def _svg_data_url(svg: str) -> str:
    return "data:image/svg+xml;base64," + base64.b64encode(svg.encode("utf-8")).decode("ascii")


_SVG = '<svg xmlns="http://www.w3.org/2000/svg" width="200" height="50"><rect width="200" height="50" fill="#123456"/></svg>'


@pytest.fixture
def use_profile(monkeypatch: pytest.MonkeyPatch):
    """Set what the three readers return; the default is an empty workspace."""

    def _set(
        profile: dict[str, str] | None = None,
        *,
        branding: dict[str, Any] | None = None,
        appearance: dict[str, Any] | None = None,
    ) -> None:
        full = dict(DEFAULT_COMPANY_PROFILE)
        full.update(profile or {})
        monkeypatch.setattr(pdf_branding, "_read_company_profile", lambda *_args, **_kw: dict(full))
        monkeypatch.setattr(pdf_branding, "_read_branding", lambda *_args, **_kw: dict(branding or {}))
        monkeypatch.setattr(pdf_branding, "_read_appearance", lambda *_args, **_kw: dict(appearance or {}))

    _set()
    return _set


def _sample(*, widths: bool = True) -> tuple[Workbook, Any]:
    """A small log sheet using every row-bound feature an exporter sets."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Log"
    for col, header in enumerate(("Ref", "Description", "Amount"), 1):
        ws.cell(row=1, column=col, value=header).font = Font(bold=True)
    for row, values in enumerate((("A-1", "Excavation", 100), ("A-2", "Concrete", 250), ("A-3", "Rebar", 75)), 2):
        for col, value in enumerate(values, 1):
            ws.cell(row=row, column=col, value=value)
    ws["B5"] = "Total"
    ws["C5"] = "=SUM(C2:C4)"
    ws["A6"] = "Note"
    ws.merge_cells("A6:C6")
    ws["A2"].hyperlink = "https://example.com/a-1"
    ws.row_dimensions[1].height = 24
    if widths:
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 40
        ws.column_dimensions["C"].width = 14
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:C4"
    ws.print_title_rows = "1:1"
    return wb, ws


def _save(wb: Workbook) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _parts(blob: bytes) -> dict[str, bytes]:
    """Every part of the package but the one carrying save timestamps."""
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        return {name: archive.read(name) for name in archive.namelist() if name != "docProps/core.xml"}


def _letterhead_values(ws: Any, start: int) -> list[Any]:
    return [cell.value for row in ws.iter_rows(min_row=1, max_row=start - 1) for cell in row if cell.value]


# -- No letterhead: nothing changes --------------------------------------------


def test_no_profile_leaves_the_workbook_identical(use_profile) -> None:
    untouched, _ = _sample()
    branded, ws = _sample()

    assert apply_company_header(ws, title="Log") == 1
    assert _parts(_save(branded)) == _parts(_save(untouched))


def test_contact_details_without_a_name_or_logo_are_no_letterhead(use_profile) -> None:
    use_profile({"address": "Hauptstrasse 1", "phone": "+49 30 1"})
    untouched, _ = _sample()
    branded, ws = _sample()

    assert apply_company_header(ws, title="Log") == 1
    assert _parts(_save(branded)) == _parts(_save(untouched))


def test_letterhead_switched_off_in_the_appearance_changes_nothing(use_profile) -> None:
    use_profile(_PROFILE, appearance={"show_letterhead": False})
    untouched, _ = _sample()
    branded, ws = _sample()

    assert apply_company_header(ws, title="Log") == 1
    assert _parts(_save(branded)) == _parts(_save(untouched))


# -- The letterhead ------------------------------------------------------------


def test_profile_prints_the_company_block_above_the_table(use_profile) -> None:
    use_profile(_PROFILE)
    wb, ws = _sample()

    start = apply_company_header(ws, title="RFI Log", subtitle="Project Alpha")
    ws = load_workbook(io.BytesIO(_save(wb)))["Log"]

    assert start > 1
    top = _letterhead_values(ws, start)
    assert top[0] == "Acme & Sons Construction GmbH"
    assert "Hauptstrasse 1 · 10115 Berlin · Germany" in top
    assert "HRB 12345 · VAT DE123456789" in top
    assert "+49 30 1234567 · office@acme.example · acme.example" in top
    assert top[-2:] == ["RFI Log", "Project Alpha"]
    # The table is intact, one block lower.
    assert [ws.cell(row=start, column=col).value for col in (1, 2, 3)] == ["Ref", "Description", "Amount"]
    assert ws.cell(row=start + 1, column=2).value == "Excavation"
    assert ws.cell(row=start, column=1).font.bold


def test_pane_filter_and_print_titles_still_point_at_the_header_row(use_profile) -> None:
    use_profile(_PROFILE)
    wb, ws = _sample()

    start = apply_company_header(ws, title="Log")
    ws = load_workbook(io.BytesIO(_save(wb)))["Log"]

    assert ws.freeze_panes == f"A{start + 1}"
    assert ws.auto_filter.ref == f"A{start}:C{start + 3}"
    # Read back, openpyxl reports the rows absolute.
    assert ws.print_title_rows.replace("$", "") == f"{start}:{start}"


def test_formulas_merges_heights_and_links_move_with_the_table(use_profile) -> None:
    use_profile(_PROFILE)
    wb, ws = _sample()
    summary = wb.create_sheet("Summary")
    summary["A1"] = "='Log'!C5"
    summary["A2"] = "=C5"

    start = apply_company_header(ws, title="Log")
    loaded = load_workbook(io.BytesIO(_save(wb)))
    ws = loaded["Log"]

    shift = start - 1
    assert ws.cell(row=5 + shift, column=3).value == f"=SUM(C{2 + shift}:C{4 + shift})"
    assert f"A{6 + shift}:C{6 + shift}" in {str(r) for r in ws.merged_cells.ranges}
    assert ws.row_dimensions[start].height == 24
    link = ws.cell(row=2 + shift, column=1).hyperlink
    assert link is not None and link.target == "https://example.com/a-1"
    # Another sheet's reference to the moved one follows it; its own does not.
    assert loaded["Summary"]["A1"].value == f"='Log'!C{5 + shift}"
    assert loaded["Summary"]["A2"].value == "=C5"


def test_print_header_and_footer_carry_the_company_name(use_profile) -> None:
    use_profile(_PROFILE)
    wb, ws = _sample()

    apply_company_header(ws, title="RFI Log")
    ws = load_workbook(io.BytesIO(_save(wb)))["Log"]

    # A literal ampersand is doubled: a single one starts a header code.
    assert ws.oddHeader.left.text == "Acme && Sons Construction GmbH"
    assert ws.oddHeader.right.text == "RFI Log"
    assert ws.oddFooter.left.text == "Acme && Sons Construction GmbH"
    assert ws.oddFooter.right.text == "&P / &N"
    # Page one has the letterhead instead of a header, and keeps the footer.
    assert ws.HeaderFooter.differentFirst
    assert not ws.firstHeader.left.text
    assert ws.firstFooter.right.text == "&P / &N"


def test_an_exporters_own_header_is_left_alone(use_profile) -> None:
    use_profile(_PROFILE)
    wb, ws = _sample()
    ws.oddHeader.center.text = "Confidential"

    apply_company_header(ws, title="Log")

    assert ws.oddHeader.center.text == "Confidential"
    assert not ws.oddHeader.left.text
    assert not ws.oddFooter.right.text


def test_the_company_name_cannot_become_a_formula(use_profile) -> None:
    use_profile({"legal_name": '=HYPERLINK("http://evil.example","x")'})
    wb, ws = _sample()

    start = apply_company_header(ws, title="=1+1")
    blob = _save(wb)
    ws = load_workbook(io.BytesIO(blob))["Log"]

    assert _letterhead_values(ws, start) == ['=HYPERLINK("http://evil.example","x")', "=1+1"]
    assert all(cell.data_type != "f" for row in ws.iter_rows(max_row=start - 1) for cell in row)
    # The table's own SUM is the only formula in the file.
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        assert archive.read("xl/worksheets/sheet1.xml").count(b"<f>") == 1


def test_control_characters_in_the_title_are_dropped(use_profile) -> None:
    use_profile({"legal_name": "Acme"})
    wb, ws = _sample()

    start = apply_company_header(ws, title="Tender\x01 BOQ\x1f")
    ws = load_workbook(io.BytesIO(_save(wb)))["Log"]

    assert _letterhead_values(ws, start) == ["Acme", "Tender BOQ"]


def test_a_phone_number_prints_as_written(use_profile) -> None:
    # neutralise_formula would print this with a leading apostrophe.
    use_profile({"legal_name": "Acme", "phone": "+1 555 0100"})
    wb, ws = _sample()

    start = apply_company_header(ws)
    ws = load_workbook(io.BytesIO(_save(wb)))["Log"]

    assert _letterhead_values(ws, start) == ["Acme", "+1 555 0100"]


# -- The logo ------------------------------------------------------------------


def _anchor_cell(picture: Any) -> tuple[int, int]:
    """(row, column), 1-based: a coordinate before saving, a marker after loading."""
    if isinstance(picture.anchor, str):
        from openpyxl.utils.cell import coordinate_to_tuple

        return coordinate_to_tuple(picture.anchor)
    marker = picture.anchor._from
    return marker.row + 1, marker.col + 1


def test_logo_is_anchored_at_the_top_left(use_profile) -> None:
    use_profile({**_PROFILE, "document_logo_data_url": _png_data_url()})
    wb, ws = _sample()

    start = apply_company_header(ws, title="Log")
    blob = _save(wb)
    ws = load_workbook(io.BytesIO(blob))["Log"]

    assert len(ws._images) == 1
    assert _anchor_cell(ws._images[0]) == (1, 1)
    assert "Acme & Sons Construction GmbH" in _letterhead_values(ws, start)
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        assert [name for name in archive.namelist() if name.startswith("xl/media/")] == ["xl/media/image1.png"]


def test_logo_fits_its_box_and_the_rows_are_tall_enough_for_it(use_profile) -> None:
    use_profile({**_PROFILE, "document_logo_data_url": _png_data_url(400, 100)})
    wb, ws = _sample(widths=False)

    start = apply_company_header(ws, title="Log")

    picture = ws._images[0]
    assert (picture.width, picture.height) == (180, 45)
    # Default columns are narrow enough for the text to sit beside the logo,
    # starting in the first column clear of it.
    beside = [cell for row in ws.iter_rows(max_row=start - 1) for cell in row if cell.value == _PROFILE["legal_name"]]
    assert beside[0].column == 4 and beside[0].row == 1
    block_pt = sum(ws.row_dimensions[row].height or 15 for row in range(1, 5))
    assert block_pt >= picture.height * 0.75


def test_text_beside_the_logo_may_start_after_a_wide_column(use_profile) -> None:
    # Column A is 75 px and B 285 px: the text starts in C, 360 px in, which
    # is still on screen when the file opens.
    use_profile({**_PROFILE, "document_logo_data_url": _png_data_url(400, 100)})
    wb, ws = _sample()

    start = apply_company_header(ws, title="Log")

    name = next(cell for row in ws.iter_rows(max_row=start - 1) for cell in row if cell.value == _PROFILE["legal_name"])
    assert (name.row, name.column) == (1, 3)


def test_wide_first_columns_put_the_text_under_the_logo(use_profile) -> None:
    use_profile({**_PROFILE, "document_logo_data_url": _png_data_url(400, 100)})
    wb, ws = _sample()
    ws.column_dimensions["B"].width = 80  # column C would start 640 px in

    start = apply_company_header(ws, title="Log")

    name = next(cell for row in ws.iter_rows(max_row=start - 1) for cell in row if cell.value == _PROFILE["legal_name"])
    assert (name.row, name.column) == (2, 1)
    assert _anchor_cell(ws._images[0]) == (1, 1)
    assert ws.row_dimensions[1].height >= ws._images[0].height * 0.75


def test_the_app_logo_stands_in_when_the_document_logo_will_not_decode(use_profile) -> None:
    use_profile(
        {**_PROFILE, "document_logo_data_url": "data:image/png;base64,bm90IGFuIGltYWdl"},
        branding={"logo_data_url": _png_data_url(64, 64)},
    )
    wb, ws = _sample()

    apply_company_header(ws, title="Log")

    assert len(ws._images) == 1
    assert (ws._images[0].width, ws._images[0].height) == (60, 60)


def test_svg_logo_is_rasterised(use_profile) -> None:
    use_profile({**_PROFILE, "document_logo_data_url": _svg_data_url(_SVG)})
    wb, ws = _sample()

    apply_company_header(ws, title="Log")
    blob = _save(wb)

    assert len(ws._images) == 1
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        media = [name for name in archive.namelist() if name.startswith("xl/media/")]
        assert media == ["xl/media/image1.png"]
        assert archive.read(media[0])[:8] == b"\x89PNG\r\n\x1a\n"


def test_broken_svg_logo_is_skipped_and_the_text_still_prints(use_profile) -> None:
    use_profile({**_PROFILE, "document_logo_data_url": _svg_data_url("<svg><not closed")})
    wb, ws = _sample()

    start = apply_company_header(ws, title="Log")
    ws = load_workbook(io.BytesIO(_save(wb)))["Log"]

    assert ws._images == []
    assert _letterhead_values(ws, start)[0] == _PROFILE["legal_name"]


def test_webp_logo_is_stored_as_png(use_profile) -> None:
    use_profile({**_PROFILE, "document_logo_data_url": _png_data_url(fmt="WEBP")})
    wb, ws = _sample()

    apply_company_header(ws, title="Log")

    with zipfile.ZipFile(io.BytesIO(_save(wb))) as archive:
        assert [name for name in archive.namelist() if name.startswith("xl/media/")] == ["xl/media/image1.png"]


def test_logo_only_profile_whose_logo_will_not_decode_adds_no_rows(use_profile) -> None:
    use_profile({"document_logo_data_url": "data:image/png;base64,bm90IGFuIGltYWdl"})
    wb, ws = _sample()

    assert apply_company_header(ws, title="Log") == 1
    assert ws["A1"].value == "Ref"
    assert ws.freeze_panes == "A2"


# -- Degrading -----------------------------------------------------------------


def test_a_sheet_with_a_chart_gets_the_print_header_but_no_rows(use_profile) -> None:
    use_profile(_PROFILE)
    wb, ws = _sample()
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=3, min_row=1, max_row=4), titles_from_data=True)
    ws.add_chart(chart, "E2")

    assert apply_company_header(ws, title="Log") == 1
    assert ws["A1"].value == "Ref"
    assert ws.oddFooter.left.text == "Acme && Sons Construction GmbH"
    assert not ws.HeaderFooter.differentFirst


def test_a_failure_leaves_the_sheet_as_it_was(use_profile, monkeypatch: pytest.MonkeyPatch) -> None:
    use_profile(_PROFILE)
    untouched, _ = _sample()
    branded, ws = _sample()

    def _boom(*_args: Any, **_kwargs: Any) -> Any:
        raise RuntimeError("layout failed")

    monkeypatch.setattr(xlsx_branding, "_layout", _boom)

    assert apply_company_header(ws, title="Log") == 1
    assert _parts(_save(branded)) == _parts(_save(untouched))


def test_write_only_sheets_are_left_alone(use_profile) -> None:
    use_profile(_PROFILE)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Catalogue")
    ws.append(["Code", "Description"])

    assert apply_company_header(ws, title="Catalogue") == 1
    # Saved so the sheet's row writer is closed here, not by the collector
    # in the middle of another test.
    _save(wb)


# -- Wired exporters -----------------------------------------------------------


def _boq_sheet() -> tuple[Workbook, Any]:
    """The first columns of the BOQ export, as the router writes them."""
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    headers = ["Pos.", "Description", "Unit", "Quantity", "Unit Rate", "Total", "Currency", "Position ID"]
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    ws.append(["01", "Earthworks", "section", None, None, None, None, "00000000-0000-0000-0000-000000000001"])
    ws.append(["01.001", "Excavation", "m3", 120, 18.5, 2220, "EUR", "00000000-0000-0000-0000-000000000002"])
    ws.append(["01.002", "Backfill", "m3", 80, 9.0, 720, "EUR", "00000000-0000-0000-0000-000000000003"])
    ws.freeze_panes = "A2"
    return wb, ws


def _both_boq_parsers() -> list[Any]:
    from app.modules.boq import router as boq_router
    from app.modules.boq.importers import excel as boq_excel

    return [boq_excel._parse_rows_from_excel, boq_router._parse_rows_from_excel]


@pytest.mark.parametrize("with_logo", [False, True])
def test_boq_export_with_a_letterhead_imports_back(use_profile, with_logo: bool) -> None:
    logo = {"document_logo_data_url": _png_data_url()} if with_logo else {}
    use_profile({**_PROFILE, **logo})
    wb, ws = _boq_sheet()
    assert apply_company_header(ws, title="Tender BOQ") > 1
    blob = _save(wb)

    for parse in _both_boq_parsers():
        rows, meta = parse(blob)
        assert [row.get("description") for row in rows] == ["Earthworks", "Excavation", "Backfill"]
        assert meta["original_columns"][:2] == ["Pos.", "Description"]


def test_boq_logo_only_letterhead_imports_back(use_profile) -> None:
    # No legal name: the first row holds nothing but the logo.
    use_profile({"document_logo_data_url": _png_data_url()})
    wb, ws = _boq_sheet()
    assert apply_company_header(ws, title=None) > 1
    blob = _save(wb)

    for parse in _both_boq_parsers():
        rows, _meta = parse(blob)
        assert [row.get("description") for row in rows] == ["Earthworks", "Excavation", "Backfill"]


@pytest.mark.parametrize("legal_name", ["Total", "Position", "Item"])
def test_boq_letterhead_whose_name_is_a_column_name_imports_back(use_profile, legal_name: str) -> None:
    # A one-word firm name equal to a header alias puts one "known column" in
    # the letterhead's first row; the table below names more, and wins.
    use_profile({"legal_name": legal_name})
    wb, ws = _boq_sheet()
    apply_company_header(ws, title="Tender BOQ")
    blob = _save(wb)

    for parse in _both_boq_parsers():
        rows, meta = parse(blob)
        assert [row.get("description") for row in rows] == ["Earthworks", "Excavation", "Backfill"]
        assert meta["original_columns"][:2] == ["Pos.", "Description"]


def test_boq_header_on_row_one_is_read_exactly_as_before(use_profile) -> None:
    # Row 1 names two known columns, so it is the header even though a row
    # under it would qualify too: files that imported before import the same way.
    wb = Workbook()
    ws = wb.active
    ws.append(["Description", "Unit"])
    ws.append(["Pos.", "Description"])
    ws.append(["Concrete", "m3"])
    blob = _save(wb)

    for parse in _both_boq_parsers():
        rows, meta = parse(blob)
        assert meta["original_columns"] == ["Description", "Unit"]
        assert [row.get("description") for row in rows] == ["Pos.", "Concrete"]


def test_boq_single_known_header_with_nothing_better_below_keeps_row_one(use_profile) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["Description", "Notes"])
    ws.append(["Concrete", "C30/37"])
    ws.append(["Formwork", "smooth"])
    blob = _save(wb)

    for parse in _both_boq_parsers():
        rows, meta = parse(blob)
        assert meta["original_columns"] == ["Description", "Notes"]
        assert [row.get("description") for row in rows] == ["Concrete", "Formwork"]


def test_boq_sheet_without_any_known_header_keeps_row_one(use_profile) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["Foo", "Bar"])
    ws.append(["one", "two"])
    blob = _save(wb)

    for parse in _both_boq_parsers():
        rows, meta = parse(blob)
        assert meta["original_columns"] == ["Foo", "Bar"]
        assert rows == []


def _report(**overrides: Any) -> bytes:
    from app.modules.reporting.exporters import export_report

    _name, _media, blob = export_report(
        fmt="xlsx",
        report_type="summary",
        title="Quarterly cost report",
        project_name="Demo project",
        currency="EUR",
        generated_at="2026-07-13T00:00:00Z",
        template_data=None,
        data_snapshot={"summary": {"total_cost": "1234.56 EUR", "positions": 12}},
        **overrides,
    )
    return blob


def test_report_workbook_without_a_profile_is_unchanged(use_profile, monkeypatch: pytest.MonkeyPatch) -> None:
    branded = _report()
    monkeypatch.setattr(xlsx_branding, "apply_company_header", lambda ws, **_kw: 1)
    assert _parts(branded) == _parts(_report())


def test_report_workbook_carries_the_letterhead(use_profile) -> None:
    use_profile(_PROFILE)
    ws = load_workbook(io.BytesIO(_report())).active

    title_row = next(row for row in range(1, 20) if ws.cell(row=row, column=1).value == "Quarterly cost report")
    assert title_row > 1
    assert ws["A1"].value == _PROFILE["legal_name"]
    # The report froze its title row; it still does, under the letterhead.
    assert ws.freeze_panes == f"A{title_row + 1}"
