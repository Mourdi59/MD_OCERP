# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""Pure-function pins for defects found in review: masking, hosting by tag, quorum, money, export, routes."""

from __future__ import annotations

import ast
import io
from datetime import date
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pytest

from app.modules.stakeholders import export, hosting
from app.modules.stakeholders import mandates as md
from app.modules.stakeholders.catalogue import DEFAULT_FRAMEWORK_ID, FRAMEWORK_KIND, register_kind
from app.modules.stakeholders.errors import StakeholderError
from app.modules.stakeholders.evidence import placement
from app.modules.stakeholders.service import COMMITMENT_FROZEN, StakeholderService, _amount

_MODULE = Path(__file__).resolve().parents[2] / "app" / "modules" / "stakeholders"
TODAY = date(2026, 10, 12)
R = md.MandateRow


@pytest.fixture(scope="module")
def fw():
    from app.modules.project_governance.catalogue_kinds import get_catalogue_kind

    register_kind()
    return get_catalogue_kind(FRAMEWORK_KIND).get_shipped(DEFAULT_FRAMEWORK_ID)


# ── rule 17: what a masked reader gets ────────────────────────────────────


def test_masked_fields_blank_the_name_carrier_and_the_address() -> None:
    row = SimpleNamespace(
        id="i",
        project_id="p",
        code="S-001",
        display_name="Jane Private",
        category="neighbour",
        contact_id=None,
        is_group=False,
        group_size_estimate=None,
        parent_stakeholder_id=None,
        user_id=None,
        interests="",
        expectations="",
        concerns="",
        relationship_owner_user_id=None,
        geo={"address": "1 Garden Lane"},
        is_confidential=True,
        status="active",
        risk_id=None,
    )
    view = {**StakeholderService.stakeholder_out(row), **StakeholderService.masked_fields()}
    assert "Jane Private" not in str(view)
    assert "Garden Lane" not in str(view)


# ── hosting by catalogue tag, never by code ───────────────────────────────


def _catalogue(area_hosts: dict[str, tuple[str, ...]], doc_hosts: dict[str, tuple[str, ...]]):
    return SimpleNamespace(
        areas=[SimpleNamespace(code=c, hosts=h) for c, h in area_hosts.items()],
        document_types=[SimpleNamespace(code=c, hosts=h) for c, h in doc_hosts.items()],
    )


def test_the_hosting_area_is_the_tagged_one_whatever_it_is_called() -> None:
    cat = _catalogue({"alpha": (), "people_and_mandates": ("stakeholders",)}, {})
    assert hosting.hosting_area_code(cat) == "people_and_mandates"


def test_a_catalogue_without_the_tag_hosts_nothing() -> None:
    untagged = SimpleNamespace(areas=[SimpleNamespace(code="alpha")], document_types=[SimpleNamespace(code="d")])
    assert hosting.hosting_area_code(untagged) is None
    assert hosting.frame_document_types(untagged, hosting.STAKEHOLDERS_TAG) == ()


def test_each_export_is_filed_under_its_own_document_tag() -> None:
    cat = _catalogue(
        {"a": ("stakeholders",)},
        {"plan": ("stakeholders",), "org": ("organisation",), "mandates": ("organisation",), "phases": ()},
    )
    assert hosting.frame_document_types(cat, hosting.EXPORT_TAGS["register"]) == ("plan",)
    assert hosting.frame_document_types(cat, hosting.EXPORT_TAGS["organisation"]) == ("org", "mandates")


def test_no_module_file_matches_a_governance_area_by_code() -> None:
    offenders = [f.name for f in _MODULE.glob("*.py") if "m01_project_management" in f.read_text(encoding="utf-8")]
    assert offenders == []


def test_evidence_without_a_host_is_reported_once_and_never_dropped() -> None:
    assert placement("a", "a", None) == "items"
    assert placement("b", "a", None) == "nothing"
    # No area hosts the module: one unknown item, on the first area only.
    assert placement("a", None, "a") == "unplaced"
    assert placement("b", None, "a") == "nothing"
    # No governance model at all: there is no gate to report to.
    assert placement("a", None, None) == "nothing"


# ── rules 11, 12, 15: who decides in a body ───────────────────────────────


def test_a_guest_without_a_vote_does_not_decide() -> None:
    body = md.Body("board", "majority", members=(md.Member(user_id="u1"), md.Member(user_id="guest", voting=False)))
    res = md.who_decides([R("x", "body", "board", "D")], "x", "", [], [body], TODAY)
    assert res.user_ids == ["guest", "u1"]
    assert res.voting_user_ids == ["u1"]


def test_an_expired_seat_does_not_raise_the_quorum() -> None:
    body = md.Body(
        "board", "unanimous", members=(md.Member(user_id="u1"), md.Member(user_id="u2", valid_to=date(2026, 10, 1)))
    )
    assert md.quorum_needed(body, TODAY) == 1
    assert "body_inquorate" not in md.who_decides([R("x", "body", "board", "D")], "x", "", [], [body], TODAY).flags
    # While the second seat was still valid, unanimity needed both.
    assert md.quorum_needed(body, date(2026, 9, 30)) == 2


def test_a_body_with_no_valid_voting_seat_is_never_quorate() -> None:
    for kind in ("majority", "unanimous", "min_members"):
        body = md.Body("board", kind, quorum_min=1, members=(md.Member(user_id="u1", valid_to=date(2026, 1, 1)),))
        res = md.who_decides([R("x", "body", "board", "D")], "x", "", [], [body], TODAY)
        assert "body_inquorate" in res.flags, kind
    empty = md.Body("board", "unanimous")
    assert "body_inquorate" in md.who_decides([R("x", "body", "board", "D")], "x", "", [], [empty], TODAY).flags


def test_a_non_voting_seat_does_not_make_a_self_approval_conflict(fw) -> None:
    body = md.Body("board", "majority", members=(md.Member(user_id="pm", voting=False), md.Member(user_id="c")))
    rows = [R("design_decision", "role", "lead_designer", "P"), R("design_decision", "body", "board", "D")]
    holders = [md.Holder("lead_designer", "pm")]
    assert md.conflicts(fw, rows, holders, [body], TODAY) == []
    voting = md.Body("board", "majority", members=(md.Member(user_id="pm"), md.Member(user_id="c")))
    assert [c.user_id for c in md.conflicts(fw, rows, holders, [voting], TODAY)] == ["pm"]


# ── money ────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("bad", ["-1", "1.00001", "100000000000000", "NaN", "abc", 1.5, True])
def test_amounts_outside_the_column_are_refused(bad) -> None:
    with pytest.raises(StakeholderError) as exc:
        _amount(bad)
    assert exc.value.code == "invalid_amount"


def test_amounts_inside_the_column_pass_as_decimal() -> None:
    assert _amount("1250.5000") == Decimal("1250.5000")
    assert _amount("99999999999999.9999") == Decimal("99999999999999.9999")
    assert _amount("") is None


def test_currency_is_frozen_with_the_amount() -> None:
    assert {"text", "cost_estimate", "currency"} <= COMMITMENT_FROZEN


# ── write routes take no client date ─────────────────────────────────────


def test_no_write_route_takes_today_from_the_caller() -> None:
    tree = ast.parse((_MODULE / "router.py").read_text(encoding="utf-8"))
    offenders = []
    for fn in tree.body:
        if not isinstance(fn, ast.AsyncFunctionDef):
            continue
        methods = {
            d.func.attr for d in fn.decorator_list if isinstance(d, ast.Call) and isinstance(d.func, ast.Attribute)
        }
        if methods & {"post", "put", "patch", "delete"} and "today" in {a.arg for a in fn.args.args}:
            offenders.append(fn.name)
    assert offenders == []


def test_no_raw_role_string_is_passed_as_a_permission() -> None:
    src = (_MODULE / "router.py").read_text(encoding="utf-8")
    assert '"role": "viewer"' not in src


# ── export: one meaning per column, one sheet per table ───────────────────


def _register_docs(fw):
    stakeholders = [
        {
            "id": "s1",
            "code": "S-001",
            "name": None,
            "masked": True,
            "anonymised": False,
            "category": "neighbour",
            "classification": "manage_closely",
            "status": "active",
            "review_due": True,
            "open_commitments": 2,
        },
        {
            "id": "s2",
            "code": "S-002",
            "name": "Road authority",
            "masked": False,
            "anonymised": False,
            "category": "permitting_authority",
            "classification": "not_assessed",
            "status": "active",
        },
    ]
    plan = [{"title": "Newsletter", "audience_stakeholder_ids": ["s1", "s2"], "status": "active", "channel": None}]
    commitments = [
        {
            "code": "C-001",
            "text": "No early deliveries",
            "stakeholder_ids": ["s1"],
            "category": "noise",
            "status": "agreed",
        }
    ]
    return export.register_documents(
        fw, stakeholders=stakeholders, plan_items=plan, commitments=commitments, project_name="P", locale="en"
    )


def _organisation_docs(fw):
    roles = {
        "holders": [{"role_code": "client_pm", "holder_name": "A. Lead", "valid_from": "2026-01-01", "valid_to": None}],
        "vacant": ["client"],
    }
    bodies = [
        {
            "code": "steering_committee",
            "name": "Steering committee",
            "members": [{"role_code": "client_pm", "valid_today": True}, {"user_id": "u", "valid_today": True}],
            "quorum_needed": 2,
            "quorum_present": 2,
        }
    ]
    matrix = {
        "rows": [
            {"decision_type": "scope_change", "holder_kind": "body", "holder_code": "steering_committee", "letter": "D"}
        ]
    }
    return export.organisation_documents(fw, roles=roles, bodies=bodies, matrix=matrix, project_name="P", locale="en")


@pytest.mark.parametrize("build", [_register_docs, _organisation_docs])
def test_every_table_has_its_own_columns_and_every_row_fits_them(fw, build) -> None:
    docs = build(fw)
    assert len(docs) == 3
    assert len({tuple(c.header for c in d.columns) for d in docs}) == 3
    for d in docs:
        headers = [c.header for c in d.columns]
        assert len(headers) == len(set(headers)), headers
        assert d.rows and all(len(r) == len(headers) for r in d.rows), d.sheet_title


def test_the_workbook_has_one_sheet_per_table_and_masks_the_private_name(fw) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(export.render(_register_docs(fw), "xlsx")))
    assert wb.sheetnames == ["Stakeholders", "Communication plan", "Commitments"]
    text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    assert "(confidential)" in text
    assert "Road authority" in text
    assert "No early deliveries" in text


def test_the_pdf_is_the_three_tables_in_one_file(fw) -> None:
    from pypdf import PdfReader

    data = export.render(_organisation_docs(fw), "pdf")
    assert data.startswith(b"%PDF")
    assert len(PdfReader(io.BytesIO(data)).pages) >= 3


# ── the daily sweep is a registered process ──────────────────────────────


def test_the_sweep_is_registered_as_a_process() -> None:
    from app.core.processes import process_registry
    from app.modules.stakeholders.jobs import PROCESS_ID, register_sweep_process

    register_sweep_process()
    register_sweep_process()
    assert process_registry.ids().count(PROCESS_ID) == 1
    assert PROCESS_ID in process_registry.processes_for_module("stakeholders")


# ── what the caller may do, computed on the server ───────────────────────


def _can(**kw):
    from app.modules.stakeholders.service import commitment_actions

    base = {
        "status": "proposed",
        "pending": False,
        "has_owner_and_due": True,
        "is_owner": False,
        "can_edit": True,
        "routed": False,
        "flags": [],
        "is_decider": True,
    }
    return commitment_actions(**{**base, **kw})


def test_every_action_carries_allowed_and_a_reason_code() -> None:
    can = _can()
    assert set(can) == {"agreed", "in_progress", "fulfilled", "verified", "withdrawn", "breached", "edit"}
    assert can["agreed"] == {"allowed": True, "reason": None}
    assert can["verified"] == {"allowed": False, "reason": "invalid_transition"}


def test_the_reason_is_the_code_the_write_would_answer() -> None:
    assert _can(is_decider=False)["agreed"]["reason"] == "not_decider"
    assert _can(is_decider=False, routed=True)["agreed"]["allowed"] is True
    assert _can(flags=["body_inquorate"])["agreed"]["reason"] == "body_inquorate"
    assert _can(flags=["mandate_missing"], routed=True)["agreed"]["reason"] == "mandate_missing"
    assert _can(has_owner_and_due=False)["agreed"]["reason"] == "owner_and_due_required"
    assert _can(status="fulfilled", is_owner=True)["verified"]["reason"] == "owner_cannot_verify"
    assert _can(status="fulfilled")["verified"]["allowed"] is True
    assert {v["reason"] for v in _can(pending=True).values()} == {"approval_pending"}
    assert {v["reason"] for v in _can(can_edit=False).values()} == {"no_permission"}
    assert not any(v["allowed"] for v in _can(status="verified").values() if v["reason"] == "invalid_transition")
