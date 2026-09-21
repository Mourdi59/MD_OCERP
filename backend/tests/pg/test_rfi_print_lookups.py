# DDC-CWICR-OE: DataDrivenConstruction · OpenConstructionERP
# Copyright (c) 2026 Artem Boiko / DataDrivenConstruction
"""The printed RFI and the printed RFI log name things; the rows only hold ids.

An RFI stores the people it names, the drawings it links and the variation
raised from it as ids. Both printable outputs resolve those ids through SQL:
the single-RFI PDF (``RFIService.generate_rfi_pdf``) and the Excel log
(``export_rfi_log``), which until this change printed three columns of raw
UUIDs where the people belong. The unit tests drive the renderer with the
names already resolved, so this file is where the lookups themselves run
against PostgreSQL:

* people come back by full name, and by email when the full name is empty;
* a linked document of another project reads as unavailable, its name never
  printed, and neither is an id that matches nothing;
* the variation is printed by code and title;
* the Excel log carries names, statuses as words, a display name that looks
  like a formula neutralised, and the print setup.
"""

from __future__ import annotations

import io
import uuid

from openpyxl import load_workbook
from pypdf import PdfReader

from app.modules.changeorders.models import ChangeOrder
from app.modules.documents.models import Document
from app.modules.projects.models import Project
from app.modules.rfi.models import RFI
from app.modules.rfi.service import RFIService
from app.modules.users.models import User


async def _seed(pg_session) -> tuple[RFI, User, User, Project]:
    tag = uuid.uuid4().hex[:8]
    raiser = User(email=f"rfi-print-r-{tag}@example.test", hashed_password="x", full_name="Maria Keller")
    # No full name, so the email is what the reader should see.
    answerer = User(email=f"rfi-print-a-{tag}@example.test", hashed_password="x", full_name="")
    pg_session.add_all([raiser, answerer])
    await pg_session.flush()

    project = Project(name="Print Test House", owner_id=raiser.id, project_code="PTH-01", currency="eur")
    other = Project(name="Another job", owner_id=raiser.id)
    pg_session.add_all([project, other])
    await pg_session.flush()

    drawing = Document(project_id=project.id, name="A-201 Elevations rev C.pdf")
    foreign = Document(project_id=other.id, name="Drawing of another project.pdf")
    variation = ChangeOrder(project_id=project.id, code="CO-003", title="Lintel change")
    pg_session.add_all([drawing, foreign, variation])
    await pg_session.flush()

    rfi = RFI(
        project_id=project.id,
        rfi_number="RFI-001",
        subject="Lintel detail at grid C/4",
        question="Which detail governs?",
        raised_by=raiser.id,
        assigned_to=answerer.id,
        ball_in_court=answerer.id,
        status="answered",
        official_response="Use the in-situ beam.",
        responded_by=answerer.id,
        responded_at="2026-09-14",
        cost_impact=True,
        cost_impact_value="12000",
        linked_drawing_ids=[str(drawing.id), str(foreign.id), str(uuid.uuid4())],
        change_order_id=str(variation.id),
    )
    pg_session.add(rfi)
    await pg_session.flush()
    return rfi, raiser, answerer, project


async def test_the_rfi_pdf_resolves_people_documents_and_the_variation(pg_session) -> None:
    rfi, raiser, answerer, _project = await _seed(pg_session)

    pdf, number = await RFIService(pg_session).generate_rfi_pdf(rfi.id, locale="en")

    assert number == "RFI-001"
    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf)).pages)
    # A long value such as an email or a UUID wraps inside its table cell, so
    # the identity checks read the text with every break taken out; a raw id
    # split over two lines would otherwise pass a "not in" check unprinted.
    flat = "".join(text.split())
    assert "Print Test House (PTH-01)" in text
    assert "Maria Keller" in text
    assert answerer.email in flat
    assert str(raiser.id) not in flat
    assert str(answerer.id) not in flat
    assert "A-201 Elevations rev C.pdf" in text
    assert "Drawingofanotherproject.pdf" not in flat
    assert "Linked documents no longer available: 2" in text
    assert "CO-003 - Lintel change" in text
    # The stored currency is lower case; the document prints the ISO code.
    assert "Yes, 12000 EUR" in text


async def test_the_rfi_log_prints_names_words_and_a_page_setup(pg_session) -> None:
    from app.modules.rfi.router import export_rfi_log

    _rfi, raiser, answerer, project = await _seed(pg_session)
    # A display name is user input, so it goes through the formula guard like
    # every other free-text cell.
    raiser.full_name = "=HYPERLINK(1)"
    await pg_session.flush()

    response = await export_rfi_log(_user=str(raiser.id), session=pg_session, project_id=project.id)
    chunks: list[bytes] = []
    async for chunk in response.body_iterator:
        chunks.append(chunk.encode() if isinstance(chunk, str) else bytes(chunk))
    sheet = load_workbook(io.BytesIO(b"".join(chunks))).active
    assert sheet is not None

    header = [cell.value for cell in sheet[1]]
    row = dict(zip(header, [cell.value for cell in sheet[2]], strict=True))
    assert row["RFI #"] == "RFI-001"
    assert row["Status"] == "Answered"
    assert row["Assigned To"] == answerer.email
    assert row["Ball-in-Court"] == answerer.email
    assert row["Raised By"] != str(raiser.id)
    assert not str(row["Raised By"]).startswith("=")
    assert "HYPERLINK" in str(row["Raised By"])

    assert sheet.page_setup.orientation == "landscape"
    assert sheet.sheet_properties.pageSetUpPr is not None
    assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.print_title_rows == "$1:$1"
